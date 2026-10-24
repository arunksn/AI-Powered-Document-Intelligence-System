"""
Regenerates the evaluator test document set in this directory.

    pip install PyMuPDF python-docx openpyxl pillow
    python sample_documents/generate_samples.py

Every supported input path is covered (digital PDF, scanned PDF, DOCX,
XLSX, JPG, PNG), and the documents are built so the cross-document checks
have something real to find when uploaded into one project -- see
sample_documents/README.md for the expected results.
"""
import os
import random

import fitz  # PyMuPDF
import openpyxl
from docx import Document as DocxDocument
from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = os.path.dirname(os.path.abspath(__file__))
random.seed(7)


def _path(name: str) -> str:
    return os.path.join(OUT, name)


# --- Text content ----------------------------------------------------------

CONTRACT_SECTIONS = [
    ("MASTER SERVICES AGREEMENT", None),
    (None, "This Master Services Agreement is entered into on January 15, 2026 between "
           "Northwind Traders LLC (\"Client\") and Acme Supplies Inc. (\"Supplier\"). "
           "The parties hereby agree to the following terms."),
    ("1. Services", "Supplier shall provide warehouse logistics services described in Schedule A."),
    ("2. Payment Terms", "Client shall pay each undisputed invoice within 30 days of receipt."),
    ("3. Term and Termination", "Either party may terminate this agreement upon 10 days written notice."),
    ("4. Limitation of Liability", "Supplier's total liability shall not exceed $50,000. "
                                   "Client's total liability shall not exceed $1,000,000."),
    ("5. Confidentiality", "Each party shall keep Confidential Information secret for a period of 3 years."),
    ("6. Intellectual Property", "All work product is assigned to Client upon payment."),
    ("7. Governing Law", "This agreement is governed by the laws of the State of Delaware."),
]

FINANCIAL_LINES = [
    "Northwind Traders LLC",
    "Income Statement and Balance Sheet",
    "For the year ended December 31, 2025",
    "",
    "Metric | FY2025 | FY2024",
    "Total Revenue | 1,250,000 | 820,000",
    "Operating Expenses | 940,000 | 700,000",
    "Gross Profit | 610,000 | 400,000",
    "Net Income | 95,000 | 88,000",
    "Total Assets | 2,100,000 | 1,900,000",
    "Total Liabilities | 2,350,000 | 1,200,000",
    "Net Cash Flow | 40,000 | 52,000",
]

NDA_LINES = [
    "MUTUAL NON-DISCLOSURE AGREEMENT",
    "",
    "This Non-Disclosure Agreement is made between",
    "Acme Supplies Inc. and Northwind Traders LLC.",
    "",
    "1. Confidential Information. The Receiving Party",
    "shall protect all Confidential Information of the",
    "Disclosing Party for a period of 2 years.",
    "",
    "2. Termination. Either party may terminate this",
    "agreement upon 30 days written notice.",
    "",
    "3. Governing Law. This agreement is governed by",
    "the laws of the State of New York.",
]

RFP_LINES = [
    "REQUEST FOR PROPOSAL - RFP 2026-04",
    "Northwind Traders LLC",
    "",
    "Scope of Work: regional freight and warehousing",
    "for three distribution centres.",
    "",
    "Proposal submission: vendors must submit a",
    "complete proposal by the bid deadline of",
    "March 31, 2026.",
    "",
    "Payment Terms: 45 days from invoice receipt.",
    "Termination: 60 days written notice.",
]

INVOICE_PHOTO_LINES = [
    "INVOICE",
    "Invoice Number: INV-1003",
    "Vendor: Acme Supplies Inc.",
    "Bill To: Northwind Traders LLC",
    "Due Date: 2026-12-01",
    "Payment terms: Net 30",
    "",
    "Pallet storage (October)     $4,200.00",
    "Forklift rental               $1,100.00",
    "Tax: $424.00",
    "Total: $5,724.00",
]


# --- Builders --------------------------------------------------------------

def build_contract_pdf():
    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for heading, body in CONTRACT_SECTIONS:
        if heading:
            size = 16 if y == 60 else 12
            page.insert_text((60, y), heading, fontsize=size, fontname="helv")
            y += size + 10
        if body:
            rect = fitz.Rect(60, y, 540, y + 60)
            page.insert_textbox(rect, body, fontsize=10, fontname="helv")
            y += 50
    doc.save(_path("01_contract_acme_msa.pdf"))


def build_invoice_docx():
    doc = DocxDocument()
    doc.add_heading("INVOICE", level=0)
    doc.add_paragraph("Invoice Number: INV-1001")
    doc.add_paragraph("Vendor: Acme Supplies Inc.")
    doc.add_paragraph("Bill To: Northwind Traders LLC")
    doc.add_paragraph("Due Date: 2026-12-15")
    doc.add_paragraph("Payment terms: Net 60")
    doc.add_heading("Line Items", level=1)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["Description", "Quantity", "Unit Price", "Amount"]):
        table.rows[0].cells[i].text = h
    for row in [
        ("Warehouse handling", "10", "250.00", "2500.00"),
        ("Pallet storage", "20", "100.00", "2000.00"),
        ("Warehouse handling", "10", "250.00", "2500.00"),  # duplicate line item
    ]:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    doc.add_paragraph("Tax: $700.00")
    doc.add_paragraph("Total: $7,700.00")
    doc.save(_path("02_invoice_acme_inv1001.docx"))


def build_invoice_xlsx():
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Prepared for Northwind Traders LLC accounts payable"
    cover["A3"] = "Notes: see LineItems sheet"

    items = wb.create_sheet("LineItems")
    rows = [
        ["Invoice Number", "INV-1002", None, None],
        ["Vendor", "Acme Supplies Inc.", None, None],
        ["Due Date", "2025-03-01", None, None],  # already in the past
        ["Payment terms", "Net 30", None, None],
        [None, None, None, None],
        ["Description", "Quantity", "Unit Price", "Amount"],
        ["Cross-dock labour", 40, 55, 2200],
        ["Pallet storage", 15, 100, 1500],
        ["Tax", None, None, 296],
        ["Total", None, None, 3996],
    ]
    for r in rows:
        items.append(r)
    wb.save(_path("03_invoice_acme_inv1002.xlsx"))


def build_financial_statement_pdf():
    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for line in FINANCIAL_LINES:
        page.insert_text((60, y), line, fontsize=11, fontname="helv")
        y += 18
    doc.save(_path("04_financial_statement_fy2025.pdf"))


def _font(size: int):
    for candidate in ("DejaVuSans.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf", "Arial.ttf"):
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def _render_page(lines: list[str], skew_degrees: float = 0.0) -> Image.Image:
    img = Image.new("L", (1700, 2200), 255)
    draw = ImageDraw.Draw(img)
    font = _font(40)
    y = 150
    for line in lines:
        draw.text((150, y), line, fill=0, font=font)
        y += 62
    if skew_degrees:
        img = img.rotate(skew_degrees, expand=False, fillcolor=255)
    return img


def build_scanned_nda_pdf():
    # Image-only PDF (no text layer) with a slight skew, so ingestion must
    # route it to OCR and the deskew step has something to correct.
    img_path = _path("_nda_tmp.png")
    _render_page(NDA_LINES, skew_degrees=1.8).convert("1").save(img_path)
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(page.rect, filename=img_path)
    doc.save(_path("05_nda_scanned.pdf"))
    os.remove(img_path)


def build_invoice_photo_jpg():
    _render_page(INVOICE_PHOTO_LINES, skew_degrees=-1.2).convert("RGB").save(
        _path("06_invoice_acme_inv1003_photo.jpg"), quality=90,
    )


def build_rfp_png():
    _render_page(RFP_LINES).save(_path("07_rfp_freight.png"))


def build_low_quality_scan_jpg():
    # Deliberately degraded: tiny, blurred, noisy. Should trip the OCR
    # confidence threshold and show the low-quality warning.
    img = _render_page(INVOICE_PHOTO_LINES, skew_degrees=4).resize((425, 550))
    img = img.filter(ImageFilter.GaussianBlur(1.6))
    px = img.load()
    for _ in range(25000):
        x, y = random.randrange(img.width), random.randrange(img.height)
        px[x, y] = random.choice((0, 255))
    img.convert("RGB").save(_path("08_low_quality_scan.jpg"), quality=35)


if __name__ == "__main__":
    build_contract_pdf()
    build_invoice_docx()
    build_invoice_xlsx()
    build_financial_statement_pdf()
    build_scanned_nda_pdf()
    build_invoice_photo_jpg()
    build_rfp_png()
    build_low_quality_scan_jpg()
    print("Sample documents written to", OUT)
